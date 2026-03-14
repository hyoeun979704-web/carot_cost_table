#!/usr/bin/env python3
"""
이미지 → Excel 자동화 스크립트
이미지(통신사 단말기 비용표 등)를 Claude Vision API로 분석하여 Excel로 변환합니다.

사용법:
    python image_to_excel.py <이미지파일1> [이미지파일2 ...] [-o 출력파일.xlsx]

예시:
    python image_to_excel.py skt_table.png kt_table.jpg -o result.xlsx
"""

import anthropic
import base64
import json
import sys
import os
import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, PatternFill
)
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────
# 1. 이미지 분석
# ─────────────────────────────────────────────

SYSTEM_PROMPT = """당신은 통신사 단말기 비용표(가격표) 이미지를 분석하여 JSON으로 변환하는 전문가입니다.
이미지에 포함된 표, 텍스트, 숫자를 빠짐없이 추출하세요.
"""

ANALYSIS_PROMPT = """이 이미지의 모든 표와 텍스트를 분석하여 아래 JSON 형식으로 반환하세요.

반환 형식:
{
  "title": "표 제목 또는 이미지 설명",
  "sections": [
    {
      "section_name": "섹션명 (예: SKT 5G 요금제, 중고 정책 등)",
      "headers": ["열1", "열2", "열3", ...],
      "rows": [
        ["값1", "값2", "값3", ...],
        ...
      ],
      "notes": ["주석1", "주석2"]
    }
  ],
  "additional_text": ["표 밖의 중요 텍스트"]
}

주의사항:
- 셀 병합이 있는 경우 해당 값을 첫 번째 셀에만 넣고 나머지는 빈 문자열("")로 채우세요.
- 숫자는 그대로 문자열로 추출하세요 (콤마, 단위 포함).
- 색상이나 강조 표시된 중요 셀은 값 앞에 "[★]" 를 붙여주세요.
- 표가 여러 개 있으면 sections 배열에 각각 추가하세요.
- JSON 외의 다른 텍스트는 절대 출력하지 마세요.
"""


def encode_image(image_path: str) -> tuple[str, str]:
    """이미지를 base64로 인코딩하고 MIME 타입을 반환합니다."""
    path = Path(image_path)
    ext = path.suffix.lower()

    mime_map = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".gif": "image/gif",
        ".webp": "image/webp",
    }
    media_type = mime_map.get(ext, "image/jpeg")

    with open(image_path, "rb") as f:
        data = base64.standard_b64encode(f.read()).decode("utf-8")

    return data, media_type


def analyze_image(client: anthropic.Anthropic, image_path: str) -> dict:
    """Claude API로 이미지를 분석하고 구조화된 데이터를 반환합니다."""
    print(f"  → 분석 중: {image_path}")

    image_data, media_type = encode_image(image_path)

    with client.messages.stream(
        model="claude-opus-4-6",
        max_tokens=8096,
        system=SYSTEM_PROMPT,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": media_type,
                            "data": image_data,
                        },
                    },
                    {"type": "text", "text": ANALYSIS_PROMPT},
                ],
            }
        ],
    ) as stream:
        response_text = stream.get_final_message().content[0].text

    # JSON 파싱
    # 응답이 ```json ... ``` 블록으로 감싸진 경우 처리
    text = response_text.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        text = "\n".join(lines[1:-1])

    try:
        return json.loads(text)
    except json.JSONDecodeError as e:
        print(f"    ⚠ JSON 파싱 오류: {e}")
        print(f"    원본 응답 (처음 500자): {text[:500]}")
        # 파싱 실패 시 raw 텍스트를 단일 섹션으로 감싸서 반환
        return {
            "title": Path(image_path).stem,
            "sections": [
                {
                    "section_name": "원본 텍스트",
                    "headers": ["내용"],
                    "rows": [[line] for line in text.split("\n") if line.strip()],
                    "notes": [],
                }
            ],
            "additional_text": [],
        }


# ─────────────────────────────────────────────
# 2. Excel 생성
# ─────────────────────────────────────────────

# 스타일 정의
HEADER_FILL = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
SECTION_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
TITLE_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

MEDIUM_BORDER = Border(
    left=Side(style="medium"),
    right=Side(style="medium"),
    top=Side(style="medium"),
    bottom=Side(style="medium"),
)


def style_cell(cell, bold=False, color="000000", bg_fill=None,
               align="left", wrap=True, border=THIN_BORDER):
    cell.font = Font(bold=bold, color=color, name="맑은 고딕", size=9)
    cell.alignment = Alignment(
        horizontal=align, vertical="center", wrap_text=wrap
    )
    if bg_fill:
        cell.fill = bg_fill
    if border:
        cell.border = border


def write_section_to_sheet(ws, section: dict, start_row: int) -> int:
    """하나의 섹션을 워크시트에 쓰고 다음 시작 행을 반환합니다."""
    row = start_row

    # 섹션 제목
    section_name = section.get("section_name", "")
    if section_name:
        cell = ws.cell(row=row, column=1, value=section_name)
        style_cell(cell, bold=True, color="FFFFFF", bg_fill=SECTION_FILL,
                   align="center", border=MEDIUM_BORDER)
        headers = section.get("headers", [])
        if len(headers) > 1:
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row, end_column=max(len(headers), 1)
            )
        row += 1

    # 헤더 행
    headers = section.get("headers", [])
    if headers:
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            style_cell(cell, bold=True, color="FFFFFF", bg_fill=HEADER_FILL,
                       align="center")
        row += 1

    # 데이터 행
    data_rows = section.get("rows", [])
    for r_idx, data_row in enumerate(data_rows):
        fill = ALT_ROW_FILL if r_idx % 2 == 0 else None
        for col_idx, value in enumerate(data_row, start=1):
            cell_fill = fill
            cell_value = str(value) if value is not None else ""

            # ★ 강조 표시 처리
            if cell_value.startswith("[★]"):
                cell_value = cell_value[3:].strip()
                cell_fill = HIGHLIGHT_FILL

            cell = ws.cell(row=row, column=col_idx, value=cell_value)
            style_cell(cell, bg_fill=cell_fill)
        row += 1

    # 주석
    notes = section.get("notes", [])
    if notes:
        row += 1  # 빈 행
        for note in notes:
            cell = ws.cell(row=row, column=1, value=f"※ {note}")
            style_cell(cell, color="666666", border=None)
            if len(headers) > 1:
                ws.merge_cells(
                    start_row=row, start_column=1,
                    end_row=row, end_column=max(len(headers), 1)
                )
            row += 1

    return row + 1  # 섹션 간 빈 행


def create_excel(all_results: list[dict], output_path: str):
    """분석 결과를 Excel 파일로 저장합니다."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 기본 시트 제거

    for result in all_results:
        image_name = result.get("image_name", "Sheet")
        title = result.get("title", image_name)
        data = result.get("data", {})

        # 시트 이름은 31자 제한
        sheet_name = Path(image_name).stem[:31]
        ws = wb.create_sheet(title=sheet_name)

        # 시트 기본 설정
        ws.sheet_view.showGridLines = True
        ws.freeze_panes = "A2"

        current_row = 1

        # 이미지 제목 행
        title_cell = ws.cell(row=current_row, column=1, value=title)
        style_cell(title_cell, bold=True, color="FFFFFF", bg_fill=TITLE_FILL,
                   align="center", border=MEDIUM_BORDER)

        # 제목 열 병합 (최대 20열까지)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=20)
        ws.row_dimensions[1].height = 25
        current_row = 2

        # 섹션별 데이터 쓰기
        sections = data.get("sections", [])
        for section in sections:
            current_row = write_section_to_sheet(ws, section, current_row)

        # 추가 텍스트
        additional = data.get("additional_text", [])
        if additional:
            current_row += 1
            header_cell = ws.cell(row=current_row, column=1, value="[기타 정보]")
            style_cell(header_cell, bold=True, color="FFFFFF",
                       bg_fill=SECTION_FILL, align="center")
            ws.merge_cells(
                start_row=current_row, start_column=1,
                end_row=current_row, end_column=5
            )
            current_row += 1
            for text in additional:
                cell = ws.cell(row=current_row, column=1, value=text)
                style_cell(cell, border=None)
                ws.merge_cells(
                    start_row=current_row, start_column=1,
                    end_row=current_row, end_column=5
                )
                current_row += 1

        # 열 너비 자동 조정
        _auto_fit_columns(ws)

    wb.save(output_path)
    print(f"\n✅ Excel 저장 완료: {output_path}")


def _auto_fit_columns(ws, min_width=8, max_width=40):
    """열 너비를 내용에 맞게 자동 조정합니다."""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                # 한글은 2배 너비로 계산
                text = str(cell.value)
                length = sum(2 if ord(c) > 127 else 1 for c in text)
                max_len = max(max_len, length)
        width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


# ─────────────────────────────────────────────
# 3. 메인
# ─────────────────────────────────────────────

def parse_args():
    parser = argparse.ArgumentParser(
        description="이미지의 표를 분석하여 Excel로 변환합니다."
    )
    parser.add_argument(
        "images",
        nargs="+",
        help="분석할 이미지 파일 경로 (JPG, PNG, GIF, WEBP)",
    )
    parser.add_argument(
        "-o", "--output",
        default="output.xlsx",
        help="출력 Excel 파일 경로 (기본값: output.xlsx)",
    )
    return parser.parse_args()


def main():
    args = parse_args()

    # 입력 파일 확인
    for img_path in args.images:
        if not os.path.exists(img_path):
            print(f"❌ 파일을 찾을 수 없습니다: {img_path}")
            sys.exit(1)

    # Claude 클라이언트 초기화
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("❌ ANTHROPIC_API_KEY 환경변수가 설정되지 않았습니다.")
        sys.exit(1)

    client = anthropic.Anthropic(api_key=api_key)

    print(f"📊 이미지 {len(args.images)}개를 분석합니다...\n")

    all_results = []
    for img_path in args.images:
        print(f"[{args.images.index(img_path)+1}/{len(args.images)}] {img_path}")
        data = analyze_image(client, img_path)
        all_results.append({
            "image_name": img_path,
            "title": data.get("title", Path(img_path).stem),
            "data": data,
        })
        print(f"    ✓ 완료 — 섹션 {len(data.get('sections', []))}개 추출")

    print("\n📝 Excel 생성 중...")
    create_excel(all_results, args.output)


if __name__ == "__main__":
    main()
